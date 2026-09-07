"""Hoja 6. Venta Perdida en el Excel por marca (pedido Franco 06-sep-2026)."""
import io, os
import pandas as pd
from openpyxl import load_workbook
import motor_v2, reportes_marcas


def test_reporte_marca_incluye_hoja_venta_perdida():
    fx = os.path.join(os.path.dirname(__file__), "fixtures", "base_mini.xlsx")
    res = motor_v2.run_analysis(fx)
    cob = res["cobertura"]; marca = reportes_marcas.marcas_reporte(cob)[0]
    fila = cob[cob["marca"].str.upper().str.strip() == marca.upper()].iloc[0]
    df_vp = pd.DataFrame([{"tienda": fila["tienda"], "sku": str(fila["sku"]), "descripcion": "MODELO X", "marca": marca,
                           "departamento": "DEPTO Y", "linea": "POLOS M/C", "semanas_en_quiebre": 3, "cobertura_sem": 1.5,
                           "stock_uds": 2, "vta_uds_sem": 1, "uds_max": 4.0, "neto_min": 100.0, "neto_max": 150.0,
                           "margen_max": 50.0, "stock_cd": 20, "on_order": 0, "evitable": True, "accion": "reponer / transferir"}])
    data = reportes_marcas.generar_reporte_marca(marca, cob, res.get("reposiciones"), None, None, None, corte="30.08.2026", df_vp=df_vp)
    wb = load_workbook(io.BytesIO(data), read_only=True)
    assert "6. Venta Perdida" in wb.sheetnames
    ws = wb["6. Venta Perdida"]; rows = list(ws.iter_rows(values_only=True))
    assert "venta perdida" in str(rows[0][0]).lower() and "S/ 100 – 150" in str(rows[0][0])
    header = [c for c in rows[1] if c]
    assert header[:5] == ["Tienda", "SKU", "Modelo", "Depto", "Línea"] and "Acción" in header
    # sin df_vp o marca sin combos → no rompe y no crea la hoja
    data2 = reportes_marcas.generar_reporte_marca(marca, cob, corte="30.08.2026", df_vp=df_vp.assign(marca="OTRA"))
    assert "6. Venta Perdida" not in load_workbook(io.BytesIO(data2), read_only=True).sheetnames
