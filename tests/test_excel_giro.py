"""Excel de giro: formato oficial (matriz SKU × tienda) que reciben los inventories."""
import io
import pandas as pd
from openpyxl import load_workbook
import vistas_excel


def test_hoja_giro_estructura():
    matriz = pd.DataFrame({"sku": [1, 2], "nombre": ["A", "B"], "categoria": ["POLOS", "CHOMPAS"],
                           "marca": ["X", "Y"], "stock_cd": [10, 20],
                           "Jockey Plaza": [3, 0], "Atocongo": [2, 5]})
    sust = pd.DataFrame({"SKU": [1, 2], "Tienda": ["Jockey Plaza", "Atocongo"], "A reponer (uds)": [3, 5]})
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        vistas_excel.hoja_giro(w, matriz, sust)
    wb = load_workbook(io.BytesIO(buf.getvalue()))
    assert wb.sheetnames == ["Giro", "Sustento"]
    ws = wb["Giro"]
    header = [c.value for c in ws[1]]
    assert header[:5] == ["SKU", "Producto", "Línea", "Marca", "Stock CD"]     # encabezado en fila 1, sin título
    assert header[-1] == "TOTAL" and "Jockey Plaza" in header and "Atocongo" in header
    assert [ws.cell(row=r, column=len(header)).value for r in (2, 3)] == [5, 5]  # TOTAL = suma de tiendas
    assert ws.freeze_panes == "F2" and ws.auto_filter.ref


def test_hoja_giro_respeta_total_existente():
    matriz = pd.DataFrame({"sku": [1], "Jockey Plaza": [4], "TOTAL": [4]})
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        vistas_excel.hoja_giro(w, matriz)
    wb = load_workbook(io.BytesIO(buf.getvalue()))
    assert wb.sheetnames == ["Giro"]
    assert [c.value for c in wb["Giro"][1]] == ["SKU", "Jockey Plaza", "TOTAL"]
