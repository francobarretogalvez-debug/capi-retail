"""Excel del reporte semanal al proveedor (C4): pestañas, formato y cuadre con los bloques."""
import io

import pandas as pd
import pytest
from openpyxl import load_workbook

import reporte_proveedor as rp
from test_reporte_proveedor import _alertas_sint, _cob, _rep_sint, _trans_sint, _vp_sint

HOJAS = ["Resumen", "1. Venta Cero (SKU)", "1b. Venta Cero x Tienda", "2a. Sobrestock", "2b. Transferencias tiendas", "2b. Detalle transferencias", "3. Ganadores", "Leyenda"]


@pytest.fixture(scope="module")
def wb_bl():
    bl = rp.bloques_marca("M", _cob(), _trans_sint(), _vp_sint(), None, _rep_sint(), _alertas_sint(), corte="30.08.2026", semana_iso="2026-35")
    return load_workbook(io.BytesIO(rp.excel_proveedor(bl))), bl


def _col(ws, header):
    heads = [c.value for c in ws[2]]
    i = heads.index(header) + 1
    return [ws.cell(row=r, column=i).value for r in range(3, ws.max_row + 1)]


def test_hojas_y_formato(wb_bl):
    wb, bl = wb_bl
    assert wb.sheetnames == HOJAS
    for h in HOJAS[1:7]:
        ws = wb[h]
        assert ws.freeze_panes == "A3" and "M —" in str(ws["A1"].value) and "30.08.2026" in str(ws["A1"].value)
    assert [c.value for c in wb["1. Venta Cero (SKU)"][2]][:3] == ["SKU", "Producto", "Línea"]


def test_excel_cuadra_con_bloques(wb_bl):
    wb, bl = wb_bl
    h = bl["hechos"]
    assert sum(_col(wb["1. Venta Cero (SKU)"], "Capital S/ (costo)")) == pytest.approx(h["b1"]["capital"])
    assert len(_col(wb["1. Venta Cero (SKU)"], "SKU")) == h["b1"]["n_skus"]
    assert sum(_col(wb["2a. Sobrestock"], "Capital S/ (costo)")) == pytest.approx(h["b2a"]["capital"])
    assert sum(_col(wb["2b. Transferencias tiendas"], "Uds a mover")) == h["b2b"]["uds"]
    assert sum(_col(wb["2b. Detalle transferencias"], "Uds a mover")) == h["b2b"]["uds"]      # detalle origen→destino suma lo mismo
    assert set(_col(wb["2b. Detalle transferencias"], "Tienda destino")) == {"T3", "T2"}
    assert len(_col(wb["3. Ganadores"], "SKU")) == h["b3"]["n_skus"]
    res = wb["Resumen"]
    assert _col(res, "Modelos")[:5] == [h["b1"]["n_4sem"], h["b1"]["n_paro"], h["b2a"]["n_skus"], h["b2b"]["n_skus"], h["b3"]["n_skus"]]
    assert h["b1"]["n_4sem"] + h["b1"]["n_paro"] == h["b1"]["n_skus"]
    assert _col(wb["1. Venta Cero (SKU)"], "Prioridad").count("⭐ TOP 80%") == h["b1"]["n_top"]


def test_bloques_vacios_no_rompen():
    solo_vc = _cob()
    solo_vc = solo_vc[solo_vc["sku"].isin([101, 103])]
    bl = rp.bloques_marca("M", solo_vc, corte="x")
    wb = load_workbook(io.BytesIO(rp.excel_proveedor(bl)))
    assert wb.sheetnames == HOJAS
    assert wb["2a. Sobrestock"]["A3"].value.startswith("Sin modelos")
    assert wb["3. Ganadores"]["A3"].value.startswith("Sin modelos")
