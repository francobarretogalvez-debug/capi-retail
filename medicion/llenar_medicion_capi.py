"""
llenar_medicion_capi.py — llena la hoja "Datos" de Capi_Medicion_Caso_Exito.xlsx
a partir del Excel de giro (empuje) y los snapshots semanales de Capi (parquet).

Uso:
  python llenar_medicion_capi.py --empuje "Empuje Casual 07.09.26.xlsx" \
      --semanas 2026-36 2026-35 2026-34 2026-33 \
      --base-semana "2026-36=Base al 06.09.xlsx" \
      --plantilla Capi_Medicion_Caso_Exito.xlsx --salida Capi_Medicion_W38_pre.xlsx \
      --fecha 2026-09-10

Fuente de datos (v2, 2026-09-12): snapshots de Capi en <snapshots-dir>/<semana>/
  - tienda.parquet   → grano SKU×tienda: stock_uds (stock pre) y vta_uds_sem (velocidad).
  - snapshot.parquet → grano SKU: stock_cd (stock CD pre).
  La primera semana de --semanas es la más reciente = estado PRE de stock; las N semanas
  juntas dan la velocidad (promedio de vta_uds_sem).
  Si una semana NO tiene parquet (p.ej. 2026-36 al 12-sep), --base-semana permite construirla
  EN MEMORIA desde la Base Micro con la MISMA función de ingesta de Capi
  (snapshots_engine.tienda.build_from_base). No escribe nada en snapshots/.

Reglas (mismas del LÉEME de la xlsx — NO cambian):
- Empujado  = SKU×tienda con unidades > 0 en el Excel de giro.
- Control   = tienda (de las 31 del giro) que NO recibió ese SKU pero lo tiene
              (stock > 0 o venta > 0 en la base). Tienda que nunca tuvo el SKU no es control.
- Tipo_Control = "Racionamiento" si Stock CD == TOTAL (se agotó el CD); "Umbral" si sobró CD.
- Velocidad (Prom_4sem_uds) = promedio de vta_uds_sem de las semanas entregadas.
- L_sku_sem = mediana de la cobertura POST (stock + uds recibidas) / velocidad de las
              tiendas servidas con velocidad > 0. Aproxima el nivel L del motor fair_share.
              El motor (motor_v2._reparto_fair_share) calcula L por bisección en memoria y
              NO lo persiste; recalcularlo exige a_reponer/cob_target del pipeline completo,
              así que se mantiene la mediana. Columna L_fuente lo deja explícito por fila.
- Cobertura_pre_sem = stock_pre / velocidad (vacío si velocidad = 0).

Limitación conocida de la ingesta (medida el 12-sep sobre Base al 06.09): build_from_base
solo guarda SKU×tienda con stock u on-order ≠ 0, así que una tienda que vendió y quedó en 0
no aparece esa semana y aquí cuenta como venta 0 (827 de 8.034 combos con venta esa semana).
Sesga la velocidad hacia abajo en tiendas en quiebre. No se corrige aquí (es de producción).
"""
import argparse, datetime as dt, os, re, sys
import pandas as pd
import openpyxl

COL_CD = 'Total  CD+Bodega Unid. (On-hand disponible)'   # nombre que usa el resto del script

MAPA_TIENDAS = {  # nombre en Excel de giro -> código de tienda en tienda.parquet (= prefijo Base Micro)
    'Arequipa': 'AQP', 'Atocongo': 'ATO', 'Breña': 'Breña', 'Cajamarca': 'CAJ', 'Callao': 'CALLAO',
    'Cayma': 'CAY', 'Chiclayo': 'CHIC', 'Chiclayo 2': 'CHII', 'Chimbote': 'CBT', 'Chorrillos': 'CHO',
    'Comas': 'CO', 'Huancayo': 'HYO', 'Ica': 'ICA', 'Iquitos': 'IQT', 'Jockey Plaza': 'JP',
    'Juliaca': 'JULIACA', 'Mega Plaza': 'LO', 'Miraflores': 'MIRAF', 'Outlet San Isidro': 'OSI',
    'Piura': 'PIU2', 'Plaza Norte': 'PLN', 'Primavera': 'PRIM', 'Pucallpa': 'PUCALPA I',
    'Puruchuco': 'PURU', 'Salaverry': 'SLVR', 'San Borja': 'SB', 'San Isidro': 'SI',
    'San Juan de Lurigancho': 'SJL', 'San Miguel': 'SM', 'Santa Anita': 'STA ANITA', 'Trujillo': 'TRUJ',
}
NO_TIENDA = {'SKU', 'Producto', 'Línea', 'Marca', 'Stock CD', 'TOTAL'}


def leer_empuje(path):
    e = pd.read_excel(path, sheet_name='Giro', header=1).dropna(axis=1, how='all')
    e['SKU'] = e['SKU'].astype('int64').astype(str)
    tiendas = [c for c in e.columns if c not in NO_TIENDA]
    faltan = [t for t in tiendas if t not in MAPA_TIENDAS]
    if faltan:
        raise SystemExit(f'Tiendas del giro sin mapeo a la base: {faltan}')
    return e, tiendas


def _consolidar_tienda(t, origen):
    """tienda.parquet (largo) → sin duplicados SKU×tienda. Misma regla del script v1:
    suma de numéricos, first en texto."""
    t = t.copy()
    t['sku'] = t['sku'].astype(str).str.strip().str.lstrip('0')
    dups = int(t.duplicated(['sku', 'tienda']).sum())
    if dups:
        num = [c for c in t.select_dtypes('number').columns]
        agg = {c: 'sum' for c in num}
        agg.update({c: 'first' for c in t.columns if c not in num and c not in ('sku', 'tienda')})
        t = t.groupby(['sku', 'tienda'], sort=False, as_index=False).agg(agg)
        print(f'[{origen}] {dups} filas duplicadas SKU×tienda consolidadas (suma de numéricos)')
    return t


def _consolidar_cd(cd, origen):
    """Serie stock_cd indexada por sku (str sin ceros a la izquierda), sin duplicados."""
    cd = cd.copy()
    cd.index = cd.index.astype(str).str.strip().str.lstrip('0')
    dups = int(cd.index.duplicated().sum())
    if dups:
        cd = cd.groupby(level=0, sort=False).sum()
        print(f'[{origen}] {dups} SKU duplicados en stock CD consolidados (suma)')
    return pd.to_numeric(cd, errors='coerce').fillna(0)


def leer_base(semana, snapshots_dir, base_xlsx=None, capi_dir=None):
    """Devuelve (DataFrame ancho indexado por SKU, meta). Misma interfaz de columnas que usaba
    la Base Micro: '<cod> Stk', '<cod> Unidades' y COL_CD, para no tocar el resto del script.

    Fuente: <snapshots_dir>/<semana>/tienda.parquet + snapshot.parquet.
    Si no existe tienda.parquet y se pasó base_xlsx, se construye en memoria con la ingesta
    de Capi (snapshots_engine.tienda.build_from_base). No persiste nada."""
    p_t = os.path.join(snapshots_dir, semana, 'tienda.parquet')
    p_s = os.path.join(snapshots_dir, semana, 'snapshot.parquet')
    if os.path.exists(p_t):
        t = pd.read_parquet(p_t)
        if not os.path.exists(p_s):
            raise SystemExit(f'{semana}: hay tienda.parquet pero falta snapshot.parquet (stock_cd)')
        s = pd.read_parquet(p_s)
        cd = _consolidar_cd(s.set_index('sku')['stock_cd'], f'{semana}/snapshot.parquet')
        origen = f'{semana}=parquet'
    elif base_xlsx:
        if capi_dir and capi_dir not in sys.path:
            sys.path.insert(0, capi_dir)
        from snapshots_engine import tienda as tsnap
        raw = pd.read_excel(base_xlsx)
        sem_nombre = tsnap.semana_de_nombre(os.path.basename(base_xlsx))
        if sem_nombre != semana:
            raise SystemExit(f'{base_xlsx} corresponde a {sem_nombre}, no a {semana}')
        t = tsnap.build_from_base(raw, semana)
        if COL_CD not in raw.columns:
            raise SystemExit(f'{base_xlsx}: falta la columna {COL_CD!r}')
        sku_col = 'Cód. Prod.' if 'Cód. Prod.' in raw.columns else raw.columns[3]
        cd = _consolidar_cd(raw.set_index(raw[sku_col].astype(str))[COL_CD], os.path.basename(base_xlsx))
        origen = f'{semana}=build_from_base({os.path.basename(base_xlsx)})'
    else:
        raise SystemExit(f'{semana}: no existe {p_t} y no se pasó --base-semana para esa semana')

    t = _consolidar_tienda(t, origen)
    stk = t.pivot(index='sku', columns='tienda', values='stock_uds').fillna(0)
    vta = t.pivot(index='sku', columns='tienda', values='vta_uds_sem').fillna(0)
    stk.columns = [f'{c} Stk' for c in stk.columns]
    vta.columns = [f'{c} Unidades' for c in vta.columns]
    b = stk.join(vta, how='outer').fillna(0)
    b[COL_CD] = cd.reindex(b.index).fillna(0)
    # SKUs con stock CD pero sin fila en tienda.parquet (sin stock en ninguna tienda)
    solo_cd = cd.index.difference(b.index)
    if len(solo_cd):
        extra = pd.DataFrame(0.0, index=solo_cd, columns=b.columns)
        extra[COL_CD] = cd.loc[solo_cd]
        b = pd.concat([b, extra])
    codigos = sorted(set(t['tienda']))
    meta = dict(origen=origen, filas_largo=len(t), skus=len(b), tiendas=codigos)
    print(f'[{origen}] {len(t):,} filas SKU×tienda · {len(b):,} SKU · {len(codigos)} tiendas')
    return b, meta


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--empuje', required=True)
    ap.add_argument('--semanas', nargs='+', required=True,
                    help='semanas ISO de Capi, la PRIMERA es la más reciente (stock pre). Ej: 2026-36 2026-35 2026-34 2026-33')
    ap.add_argument('--snapshots-dir', default=os.path.expanduser('~/capi-retail/snapshots'))
    ap.add_argument('--capi-dir', default=os.path.expanduser('~/capi-retail'),
                    help='raíz del repo Capi (para importar snapshots_engine si hace falta construir una semana)')
    ap.add_argument('--base-semana', action='append', default=[],
                    help='SEMANA=ruta.xlsx — Base Micro para construir en memoria una semana sin parquet')
    ap.add_argument('--plantilla', required=True)
    ap.add_argument('--salida', required=True)
    ap.add_argument('--fecha', default='2026-09-10')
    a = ap.parse_args()

    fallback = {}
    for item in a.base_semana:
        if '=' not in item:
            raise SystemExit(f'--base-semana espera SEMANA=ruta.xlsx, recibió {item!r}')
        k, v = item.split('=', 1)
        fallback[k.strip()] = os.path.expanduser(v.strip())

    e, tiendas = leer_empuje(a.empuje)
    bases, metas = [], []
    for w in a.semanas:
        b, m = leer_base(w, a.snapshots_dir, fallback.get(w), a.capi_dir)
        bases.append(b); metas.append(m)
    b0 = bases[0]  # la más reciente = stock pre
    n_bases = len(bases)
    fecha = dt.datetime.strptime(a.fecha, '%Y-%m-%d')

    # Verificar el mapeo giro → códigos reales del parquet (no asumir que coincide)
    codigos_b0 = set(metas[0]['tiendas'])
    sin_codigo = [t for t in tiendas if MAPA_TIENDAS[t] not in codigos_b0]
    if sin_codigo:
        raise SystemExit(f'Tiendas del giro cuyo código no aparece en {metas[0]["origen"]}: '
                         f'{[(t, MAPA_TIENDAS[t]) for t in sin_codigo]}')

    filas, avisos = [], []
    for _, r in e.iterrows():
        sku = r['SKU'].lstrip('0')
        if sku not in b0.index:
            avisos.append(f'SKU {sku} no está en la base'); continue
        fb = b0.loc[sku]
        tipo = 'Racionamiento' if r['Stock CD'] <= r['TOTAL'] else 'Umbral'
        cd_pre = float(fb[COL_CD])
        # velocidad y stock por tienda
        info = {}
        for t in tiendas:
            p = MAPA_TIENDAS[t]
            stk = float(fb.get(f'{p} Stk', 0) or 0)
            vel = sum(float(bb.loc[sku].get(f'{p} Unidades', 0) or 0) for bb in bases if sku in bb.index) / n_bases
            vel = max(vel, 0.0)
            uds = float(r[t] or 0)
            info[t] = dict(stk=max(stk, 0.0), vel=vel, uds=uds)
        # L del SKU: mediana cobertura post de tiendas servidas con velocidad > 0
        post = [(i['stk'] + i['uds']) / i['vel'] for i in info.values() if i['uds'] > 0 and i['vel'] > 0]
        L = float(pd.Series(post).median()) if post else None
        L_fuente = 'mediana_cob_post' if L is not None else ''
        for t, i in info.items():
            if i['uds'] > 0:
                grupo = 'Empujado'
            elif i['stk'] > 0 or i['vel'] > 0:
                grupo = 'Control'
            else:
                continue
            cob = i['stk'] / i['vel'] if i['vel'] > 0 else None
            filas.append(dict(
                ID_Par='', Tipo_Control=tipo, Grupo=grupo, Cod_Modelo=r['SKU'],
                Descripcion=r['Producto'], Marca=r['Marca'], Categoria=r['Línea'], Tienda=t,
                L_sku_sem=round(L, 2) if L is not None else None,
                Prom_4sem_uds=round(i['vel'], 2), Cobertura_pre_sem=round(cob, 2) if cob is not None else None,
                Stock_tienda_pre_uds=int(i['stk']), Stock_CD_pre_uds=int(cd_pre),
                Uds_empujadas=int(i['uds']) if i['uds'] > 0 else 0,
                Fecha_empuje=fecha if i['uds'] > 0 else None,
                L_fuente=L_fuente,
            ))

    df = pd.DataFrame(filas)
    # --- escribir en la plantilla ---
    wb = openpyxl.load_workbook(a.plantilla)
    ws = wb['Datos']
    cols = ['ID_Par', 'Tipo_Control', 'Grupo', 'Cod_Modelo', 'Descripcion', 'Marca', 'Categoria', 'Tienda',
            'L_sku_sem', 'Prom_4sem_uds', 'Cobertura_pre_sem', 'Stock_tienda_pre_uds', 'Stock_CD_pre_uds',
            'Uds_empujadas', 'Fecha_empuje']
    ws.delete_rows(4, ws.max_row)  # borra el ejemplo
    ws.cell(row=3, column=21, value='L_fuente')  # col U: cómo se obtuvo L (mediana_cob_post | motor)
    for k, row in enumerate(df.itertuples(index=False), start=4):
        for j, c in enumerate(cols, start=1):
            ws.cell(row=k, column=j, value=getattr(row, c))
        ws.cell(row=k, column=19, value=f'=IFERROR(P{k}/J{k},"")')
        ws.cell(row=k, column=20, value=f'=IFERROR(K{k}-I{k},"")')
        ws.cell(row=k, column=21, value=row.L_fuente)
    ult = 3 + len(df)
    # extender rangos de fórmulas en el resto del libro
    pat = re.compile(r'\$(1000)\b')
    for hoja in wb.worksheets:
        if hoja.title == 'Datos':
            continue
        for fila in hoja.iter_rows():
            for c in fila:
                if isinstance(c.value, str) and c.value.startswith('=') and 'Datos!' in c.value:
                    c.value = pat.sub(f'${max(ult, 1000)}', c.value)
    nota = (f'Generado {dt.date.today()} · stock pre = {metas[0]["origen"]} · velocidad = promedio de '
            f'{n_bases} semana(s): ' + ', '.join(m['origen'] for m in metas)
            + (' ⚠️ UNA sola semana: Prom_4sem es venta de 1 semana, no de 4' if n_bases == 1 else '')
            + ' · L = mediana cob post (motor no persiste L)')
    ws.cell(row=2, column=1, value=nota)
    wb.save(a.salida)

    # --- resumen ---
    print(nota)
    print(f'Filas: {len(df)} | Empujados: {(df.Grupo=="Empujado").sum()} | Control: {(df.Grupo=="Control").sum()}')
    print(df.groupby(['Tipo_Control', 'Grupo']).size())
    banda = df[(df.Cobertura_pre_sem.notna()) & ((df.Cobertura_pre_sem - df.L_sku_sem).abs() <= 1.5)]
    print(f'En banda ±1.5 sem alrededor de L: {len(banda)} filas ({(banda.Grupo=="Empujado").sum()} emp / {(banda.Grupo=="Control").sum()} ctrl)')
    print(f'Filas sin cobertura (velocidad 0): {df.Cobertura_pre_sem.isna().sum()}')
    print(f'Filas sin L (SKU sin tienda servida con velocidad > 0): {df.L_sku_sem.isna().sum()}')
    for w in avisos[:10]:
        print('AVISO:', w)


if __name__ == '__main__':
    main()
