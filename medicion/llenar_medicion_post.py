"""
llenar_medicion_post.py — llena el POST de Capi_Medicion_W38_pre.xlsx desde los snapshots de Capi.

Uso:
  python llenar_medicion_post.py --medicion Capi_Medicion_W38_pre.xlsx --semana-post 2026-38 \
      --salida Capi_Medicion_W38_post.xlsx

Qué llena (todo desde <snapshots-dir>/<semana>/, nada a mano):
- Hoja Datos, columnas P–R por fila SKU×tienda, con tienda.parquet de la semana medida:
    Venta_sem_uds     = vta_uds_sem
    Venta_sem_soles   = vta_soles_sem   (columna agregada a la ingesta el 2026-09-12)
    Quiebre_en_semana = "Sí" si el SKU×tienda está en quiebre según la regla de Capi
                        (venta_perdida_semanal, regla Franco 06-sep): cobertura en la tienda al
                        cierre ≤ 4 semanas, con velocidad de referencia = máx(prom simple, prom
                        reciente) de hasta 4 semanas previas (mín. 2 observadas) y las mismas
                        exclusiones de liquidación. Misma definición que la hoja Quiebres.
    Estado_cob_fin (col W) = Quiebre (≤4 sem) · Pre-quiebre (4–8 sem, mismo umbral que
                        taxonomia.Estado.PRE_QUIEBRE) · "Cob > 8 o sin velocidad".
    Cobertura_fin_sem (col X) = cobertura al cierre según Capi, solo para filas ≤ 8 sem.
  Una fila ausente en tienda.parquet significa stock 0 y on-order 0 al cierre (la ingesta no la
  guarda) → venta 0. OJO: si esa tienda vendió algo antes de quedar en 0, esa venta
  se pierde (filtro de ingesta, ver Pendientes-Post-Presentacion-Zina). Columna V marca esas filas.
- Hoja Quiebres: por semana disponible, combos SKU×tienda en quiebre con stock en CD y su venta
  perdida neta máx; columna G: combos en pre-quiebre (4–8 sem) con stock en CD. El umbral de 8 se
  aplica cambiando vp.COB_QUIEBRE_SEM solo durante la llamada (la función no lo recibe como
  parámetro; no se toca producción) (venta_perdida_semanal.venta_perdida_semana, mismo cálculo que reprodujo la
  referencia del LÉEME: 744 / S/ 56K en la semana 35).
- Hoja Canibalización: venta en soles de la CATEGORÍA COMPLETA (todos los SKU de esas líneas, no
  solo los empujados) en tiendas empujadas vs control. Tienda empujada de una categoría = recibió
  ≥1 unidad de algún SKU de esa categoría en el giro; control = las demás tiendas del giro.
  Venta previa = promedio semanal de las N semanas anteriores con soles en el parquet.
"""
import argparse, datetime as dt, os, sys
import pandas as pd
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from llenar_medicion_capi import MAPA_TIENDAS  # noqa: E402

COB_PRE_QUIEBRE = 8.0   # taxonomia.py: PRE-QUIEBRE = cobertura 4–8 sem (quiebre ≤ 4 = vp.COB_QUIEBRE_SEM)

CATEGORIAS = {  # fila de la hoja Canibalización -> líneas del snapshot
    'Polos': ['POLOS M/C', 'POLOS M/L'], 'Camisas': ['CAMISAS M/C', 'CAMISAS M/L'],
    'Pantalones': ['PANTALONES'], 'Shorts': ['SHORTS'], 'Chompas': ['CHOMPAS'], 'Denim': ['JEANS'],
    'Casacas': ['CASACAS'], 'Otros del giro': ['ACCESORIOS', 'POLERONES', 'BLAZERS'],
}


def _sku(s):
    return s.astype(str).str.strip().str.lstrip('0')


def leer_tienda(snapshots_dir, semana):
    p = os.path.join(snapshots_dir, semana, 'tienda.parquet')
    if not os.path.exists(p):
        raise SystemExit(f'No existe {p}. Sube la base de esa semana a Capi primero.')
    t = pd.read_parquet(p)
    if 'vta_soles_sem' not in t.columns:
        raise SystemExit(f'{p} no tiene vta_soles_sem: regenerar con la ingesta actual (rama ingesta-vta-soles)')
    t['sku'] = _sku(t['sku'])
    if t.duplicated(['sku', 'tienda']).any():
        raise SystemExit(f'{p} trae SKU×tienda duplicados; consolidar antes de medir')
    return t


def semanas_previas(snapshots_dir, semana_post, n):
    """Últimas n semanas anteriores a semana_post con tienda.parquet y soles > 0."""
    out = []
    for d in sorted(os.listdir(snapshots_dir), reverse=True):
        p = os.path.join(snapshots_dir, d, 'tienda.parquet')
        if d >= semana_post or not os.path.exists(p):
            continue
        if pd.read_parquet(p, columns=['vta_soles_sem'])['vta_soles_sem'].sum() <= 0:
            print(f'[{d}] sin soles por tienda (base formato viejo) → no entra a la venta previa')
            continue
        out.append(d)
        if len(out) == n:
            break
    return sorted(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--medicion', required=True, help='xlsx con el PRE ya llenado (hoja Datos A–O)')
    ap.add_argument('--semana-post', required=True, help='semana ISO medida, ej. 2026-38')
    ap.add_argument('--salida', required=True)
    ap.add_argument('--snapshots-dir', default=os.path.expanduser('~/capi-retail/snapshots'))
    ap.add_argument('--capi-dir', default=os.path.expanduser('~/capi-retail'))
    ap.add_argument('--n-pre', type=int, default=4, help='semanas previas para la venta previa de Canibalización')
    a = ap.parse_args()
    if a.capi_dir not in sys.path:
        sys.path.insert(0, a.capi_dir)

    wb = openpyxl.load_workbook(a.medicion)
    ws = wb['Datos']
    hdr = [c.value for c in ws[3]]
    col = {h: i + 1 for i, h in enumerate(hdr) if h}
    for h in ('Grupo', 'Cod_Modelo', 'Categoria', 'Tienda', 'Uds_empujadas', 'Venta_sem_uds', 'Venta_sem_soles', 'Quiebre_en_semana'):
        if h not in col:
            raise SystemExit(f'La hoja Datos no tiene la columna {h}')
    col_flag = max(col.values()) + 1
    col_estado, col_cob = col_flag + 1, col_flag + 2
    ws.cell(row=3, column=col_flag, value='Post_ausente_en_parquet')
    ws.cell(row=3, column=col_estado, value='Estado_cob_fin')
    ws.cell(row=3, column=col_cob, value='Cobertura_fin_sem')

    # ── 1. Datos P–R ──
    post = leer_tienda(a.snapshots_dir, a.semana_post).set_index(['sku', 'tienda'])
    import venta_perdida_semanal as vp

    def vp_con_umbral(w, cob):
        """venta_perdida_semana con otro umbral de cobertura; restaura la constante al salir."""
        prev_c = vp.COB_QUIEBRE_SEM
        vp.COB_QUIEBRE_SEM = cob
        try:
            return vp.venta_perdida_semana(w)
        finally:
            vp.COB_QUIEBRE_SEM = prev_c

    def claves(res):
        q = res['en_quiebre']
        return set(zip(_sku(q['sku']), q['tienda'])) if len(q) else set()

    res_post = vp.venta_perdida_semana(a.semana_post)
    res_post8 = vp_con_umbral(a.semana_post, COB_PRE_QUIEBRE)
    quiebres = claves(res_post)
    q8 = res_post8['en_quiebre']
    cob_fin = dict(zip(zip(_sku(q8['sku']), q8['tienda']), q8['cobertura_sem'])) if len(q8) else {}
    print(f'Regla Capi: quiebre cob ≤ {vp.COB_QUIEBRE_SEM:g} sem → {len(quiebres):,} SKU×tienda; pre-quiebre (4–{COB_PRE_QUIEBRE:g}] → '
          f'{len(cob_fin) - len(quiebres):,} en {a.semana_post}; velocidad de {res_post["prev"]}')
    filas, ausentes = [], 0
    for r in range(4, ws.max_row + 1):
        grupo = ws.cell(row=r, column=col['Grupo']).value
        if grupo not in ('Empujado', 'Control'):
            continue
        sku = str(ws.cell(row=r, column=col['Cod_Modelo']).value).strip().lstrip('0')
        tienda = ws.cell(row=r, column=col['Tienda']).value
        codigo = MAPA_TIENDAS.get(tienda)
        if codigo is None:
            raise SystemExit(f'Fila {r}: tienda {tienda!r} sin código en MAPA_TIENDAS')
        key = (sku, codigo)
        if key in post.index:
            x = post.loc[key]
            uds, soles, ausente = int(x['vta_uds_sem']), float(x['vta_soles_sem']), False
        else:
            uds, soles, ausente = 0, 0.0, True
            ausentes += 1
        quiebre = 'Sí' if key in quiebres else 'No'
        cob = cob_fin.get(key)
        estado = 'Quiebre' if key in quiebres else ('Pre-quiebre' if cob is not None else f'Cob > {COB_PRE_QUIEBRE:g} o sin velocidad')
        ws.cell(row=r, column=col_estado, value=estado)
        if cob is not None:
            ws.cell(row=r, column=col_cob, value=round(float(cob), 2))
        ws.cell(row=r, column=col['Venta_sem_uds'], value=uds)
        ws.cell(row=r, column=col['Venta_sem_soles'], value=round(soles, 2))
        ws.cell(row=r, column=col['Quiebre_en_semana'], value=quiebre)
        ws.cell(row=r, column=col_flag, value='Sí' if ausente else 'No')
        filas.append(dict(grupo=grupo, tipo=ws.cell(row=r, column=col['Tipo_Control']).value, cat=ws.cell(row=r, column=col['Categoria']).value,
                          tienda=tienda, uds=uds, soles=soles, quiebre=quiebre == 'Sí', prequiebre=estado == 'Pre-quiebre', emp=float(ws.cell(row=r, column=col['Uds_empujadas']).value or 0) > 0))
    df = pd.DataFrame(filas)
    nota_prev = ws.cell(row=2, column=1).value or ''
    ws.cell(row=2, column=1, value=f'{nota_prev} · POST {a.semana_post} desde tienda.parquet ({dt.date.today()}), {ausentes} filas ausentes en el parquet (stock 0 al cierre) · Quiebre_en_semana = regla Capi cob ≤ 4 sem')

    # ── 2. Quiebres ──
    wq = wb['Quiebres']
    wq.cell(row=3, column=7, value=f'Pre-quiebre (4–{COB_PRE_QUIEBRE:g} sem) con stock en CD')
    q_lleno = []
    for r in range(4, wq.max_row + 1):
        lab = wq.cell(row=r, column=1).value
        if not (isinstance(lab, str) and lab.startswith('W') and lab[1:].isdigit()):
            continue
        w = f'{a.semana_post[:4]}-{int(lab[1:]):02d}'
        if w > a.semana_post or not os.path.exists(os.path.join(a.snapshots_dir, w, 'tienda.parquet')):
            continue
        res = res_post if w == a.semana_post else vp.venta_perdida_semana(w)
        res8 = res_post8 if w == a.semana_post else vp_con_umbral(w, COB_PRE_QUIEBRE)
        d = res['detalle']
        cd = d[d['stock_cd'] > 0] if 'stock_cd' in d.columns else d.iloc[0:0]
        # pre-quiebre: (4, 8] sem con stock en CD (stock_cd del snapshot de cadena de esa semana)
        s_cd = pd.read_parquet(os.path.join(a.snapshots_dir, w, 'snapshot.parquet'), columns=['sku', 'stock_cd'])
        s_cd = dict(zip(_sku(s_cd['sku']), s_cd['stock_cd']))
        p8 = res8['en_quiebre']
        p8 = p8[p8['cobertura_sem'] > vp.COB_QUIEBRE_SEM] if len(p8) else p8
        n_pre_cd = int(sum(1 for k in _sku(p8['sku']) if float(s_cd.get(k, 0) or 0) > 0)) if len(p8) else 0
        wq.cell(row=r, column=3, value=int(len(cd)))
        wq.cell(row=r, column=4, value=round(float(cd['neto_max'].sum()), 0))
        wq.cell(row=r, column=6, value=f'Capi venta_perdida_semanal {w}: combos en quiebre con stock CD > 0, pérdida neta máx')
        wq.cell(row=r, column=7, value=n_pre_cd)
        q_lleno.append((lab, w, int(len(cd)), round(float(cd['neto_max'].sum())), n_pre_cd))

    # ── 3. Canibalización ──
    prev = semanas_previas(a.snapshots_dir, a.semana_post, a.n_pre)
    if not prev:
        raise SystemExit('Sin semanas previas con soles para Canibalización')
    lineas = {}
    for w in prev + [a.semana_post]:
        s = pd.read_parquet(os.path.join(a.snapshots_dir, w, 'snapshot.parquet'), columns=['sku', 'linea'])
        lineas.update(dict(zip(_sku(s['sku']), s['linea'])))
    tp = {w: leer_tienda(a.snapshots_dir, w) for w in prev}
    tp[a.semana_post] = post.reset_index()
    for w, t in tp.items():
        t['linea'] = t['sku'].map(lineas)
    codigo_giro = set(MAPA_TIENDAS.values())
    wc = wb['Canibalización']
    fila_cat = {wc.cell(row=r, column=1).value: r for r in range(4, 12)}
    libres = [r for r in range(4, 12) if not wc.cell(row=r, column=1).value]
    canib = []
    for cat, lins in CATEGORIAS.items():
        r = fila_cat.get(cat)
        if r is None:
            if not libres:
                print(f'[Canibalización] sin fila libre para {cat}; se omite'); continue
            r = libres.pop(0); wc.cell(row=r, column=1, value=cat)
        emp_t = {MAPA_TIENDAS[t] for t in df[(df.emp) & (df.cat.isin(lins))].tienda.unique()}
        ctl_t = codigo_giro - emp_t
        def venta(w, tiendas):
            t = tp[w]
            return float(t[(t['linea'].isin(lins)) & (t['tienda'].isin(tiendas))]['vta_soles_sem'].sum())
        prev_emp = sum(venta(w, emp_t) for w in prev) / len(prev)
        prev_ctl = sum(venta(w, ctl_t) for w in prev) / len(prev)
        post_emp, post_ctl = venta(a.semana_post, emp_t), venta(a.semana_post, ctl_t)
        for c, v in zip((2, 3, 5, 6), (prev_emp, post_emp, prev_ctl, post_ctl)):
            wc.cell(row=r, column=c, value=round(v, 0))
        canib.append(dict(cat=cat, n_emp=len(emp_t), n_ctl=len(ctl_t), prev_emp=round(prev_emp), post_emp=round(post_emp), prev_ctl=round(prev_ctl), post_ctl=round(post_ctl)))
        if min(len(emp_t), len(ctl_t)) < 5:
            print(f'[Canibalización] {cat}: {len(emp_t)} tiendas empujadas vs {len(ctl_t)} control → la comparación por tienda '
                  f'no sostiene (el giro cubrió casi todas las tiendas); leer solo la variación de las empujadas o comparar cadena completa')
    wc.cell(row=13, column=1, value=f'Generado {dt.date.today()} · categoría completa (todos los SKU de la línea) · venta previa = promedio de {", ".join(prev)} · semana medida {a.semana_post} · fuente tienda.parquet')

    wb.save(a.salida)

    # ── resumen ──
    print(f'POST {a.semana_post}: {len(df)} filas llenadas · {ausentes} ausentes en parquet (stock 0 al cierre, venta 0)')
    g = df.groupby(['tipo', 'grupo']).agg(n=('uds', 'size'), uds_prom=('uds', 'mean'), soles_prom=('soles', 'mean'),
                                          quiebre_pct=('quiebre', 'mean'), prequiebre_pct=('prequiebre', 'mean')).round(2)
    print(g.to_string())
    print('Quiebres (semana, ISO, quiebres con CD, S/ perdida, pre-quiebre con CD):', q_lleno if q_lleno else 'ninguna semana disponible')
    print('Canibalización (semanas previas', prev, '):'); print(pd.DataFrame(canib).to_string(index=False))


if __name__ == '__main__':
    main()
