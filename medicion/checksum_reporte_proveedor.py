#!/usr/bin/env python3
"""Checksum del reporte semanal al proveedor sobre una base REAL (C5, decisión 2026-09-18).

Uso:
    python medicion/checksum_reporte_proveedor.py --base "data2/bases antiguas/Base al 30.08.xlsx" --marca "JOHN HOLDEN"
    python medicion/checksum_reporte_proveedor.py --base "..." --marca ALL [--salida /tmp/dir]

Pasa la base por transformar_profundidad.transform (la MISMA ruta que la app; la lectura directa
con formato 'ripley' devuelve ventas en cero — medido 2026-09-19) y compara el Excel de 9 pestañas
con los bloques del proveedor. Exit 0 solo si todos los cuadres dan Δ = 0. Solo lectura.
"""
import argparse
import io
import os
import sys
import tempfile
import time
import warnings

warnings.filterwarnings("ignore")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import motor_v2  # noqa: E402
import pricing  # noqa: E402
import reporte_proveedor as rp  # noqa: E402
import reportes_marcas as rm  # noqa: E402
import transformar_profundidad as etl  # noqa: E402


def _col(ws, header):
    heads = [c.value for c in ws[2]]
    if header not in heads:
        return []
    i = heads.index(header) + 1
    return [ws.cell(row=r, column=i).value for r in range(3, ws.max_row + 1)]


def cargar(base):
    import openpyxl
    wb = openpyxl.load_workbook(base, read_only=True); sheets = wb.sheetnames; wb.close()
    if "Base" in sheets:
        pl = os.path.join(tempfile.mkdtemp(), "plantilla.xlsx")
        etl.transform(base, output_path=pl, fecha_corte=etl.fecha_corte_desde_nombre(os.path.basename(base)))
    else:
        pl = base
    return motor_v2.run_analysis(pl)


def checksum_marca(res, marca, df_vp, corte, salida=None):
    args = (marca, res["cobertura"], res["reposiciones"], res["transferencias"], res["acciones_precio"], res["alertas"], corte, df_vp)
    wb9 = load_workbook(io.BytesIO(rm.generar_reporte_marca(*args)))
    bl = rp.bloques_marca(marca, res["cobertura"], res["transferencias"], df_vp, res["acciones_precio"], res["reposiciones"], res["alertas"], corte=corte)
    xb = rp.excel_proveedor(bl); wbp = load_workbook(io.BytesIO(xb))
    h = bl["hechos"]; filas = []
    def add(nombre, a, b, ok=None, nota=""):
        ok = (abs((a or 0) - (b or 0)) < 0.5) if ok is None else ok
        filas.append({"cuadre": nombre, "9 hojas": a, "bloques": b, "OK": "✅" if ok else "❌", "nota": nota})
        return ok
    oks = []
    cap9 = sum(v or 0 for v in _col(wb9["5. Venta Cero"], "Capital S/")) if "5. Venta Cero" in wb9.sheetnames else 0
    capp = sum(v or 0 for v in _col(wbp["1b. Venta Cero x Tienda"], "Capital S/"))
    oks.append(add("Capital venta cero x tienda (hoja 5 vs 1b)", round(cap9), round(capp)))
    oks.append(add("Capital B1 hoja '1. Venta Cero (SKU)' vs hechos", round(sum(v or 0 for v in _col(wbp["1. Venta Cero (SKU)"], "Capital S/ (costo)"))), h["b1"]["capital"]))
    oks.append(add("Capital B2a hoja '2a' vs hechos", round(sum(v or 0 for v in _col(wbp["2a. Sobrestock"], "Capital S/ (costo)"))), h["b2a"]["capital"]))
    h4 = {int(k): int(v) for k, v in zip(_col(wb9["4. Transferir"], "SKU"), _col(wb9["4. Transferir"], "Uds a mover"))} if "4. Transferir" in wb9.sheetnames else {}
    b2b = {int(k): int(v) for k, v in zip(bl["b2b"]["sku"], bl["b2b"]["transf_uds"])} if not bl["b2b"].empty else {}
    oks.append(add("Uds a mover por SKU (hoja 4 vs 2b)", sum(h4.values()), sum(b2b.values()), ok=(h4 == b2b), nota=f"{len(h4)} vs {len(b2b)} SKUs"))
    ref = {}
    for hoja in ("1. Liquidar", "2. Activar"):
        if hoja in wb9.sheetnames:
            for sku, est, p in zip(_col(wb9[hoja], "SKU"), _col(wb9[hoja], "Estado"), _col(wb9[hoja], "P. Sugerido")):
                ref.setdefault(int(sku), {})[est] = p
    b2a = bl["b2a"]; n_sol = sum(int(s) in ref for s in b2a["sku"]) if not b2a.empty else 0
    maxd = 0.0
    if not b2a.empty:
        for _, r in b2a.iterrows():
            p9 = ref.get(int(r["sku"]), {}).get(r["estado_cadena"])
            if p9 is not None and pd.notna(r["precio_sugerido"]):
                maxd = max(maxd, abs(float(p9) - float(r["precio_sugerido"])))
    filas.append({"cuadre": "SKUs B2a presentes en hojas 1+2 (informativo)", "9 hojas": len(ref), "bloques": f"{n_sol}/{len(b2a)}", "OK": "ℹ️", "nota": "B2a = estado de CADENA; hojas 1+2 = estado por tienda (grano distinto a propósito)"})
    oks.append(add("P. Sugerido por SKU (max |Δ| hoja 2 vs B2a)", 0, round(maxd, 2), ok=maxd < 0.01))
    margen_min = float(motor_v2.DEFAULT_PARAMS.get("margen_min", 0.15)); viol = 0; n_p = 0
    for df in (bl["b1"], b2a):
        if df.empty or "precio_sugerido" not in df.columns:
            continue
        con = df[df["precio_sugerido"].notna()]
        n_p += len(con)
        viol += int(((con["precio_sugerido"] >= con["precio_vigente"]) | (con["precio_sugerido"] < con.apply(lambda r: pricing.precio_piso(r["costo"], margen_min) - 0.01, axis=1))).sum())
    oks.append(add("Precios sugeridos que suben o rompen piso", 0, viol, ok=(viol == 0), nota=f"{n_p} precios revisados"))
    s1, s2, s3 = set(bl["b1"]["sku"]), set(b2a["sku"]) if not b2a.empty else set(), set(bl["b3"]["sku"]) if not bl["b3"].empty else set()
    oks.append(add("SKUs en más de un bloque", 0, len(s1 & s2) + len(s1 & s3) + len(s2 & s3), ok=not (s1 & s2 or s1 & s3 or s2 & s3)))
    filas.append({"cuadre": "B3 (informativo)", "9 hojas": "", "bloques": f"{h['b3']['n_skus']} SKUs · umbral {h['b3']['umbral_vta']} u/sem · {h['b3']['n_sin_cd']} sin CD · vp {h['b3']['vp_neto_min']}–{h['b3']['vp_neto_max']}", "OK": "ℹ️", "nota": ""})
    if salida:
        os.makedirs(salida, exist_ok=True)
        open(os.path.join(salida, f"Reporte_Proveedor_{marca.replace(' ', '_')}_{corte}.xlsx"), "wb").write(xb)
        open(os.path.join(salida, f"correo_{marca.replace(' ', '_')}.txt"), "w").write("\n\n".join(f"[{k}]\n{v}" for k, v in rp.tablas_texto(bl).items()))
    return pd.DataFrame(filas), all(oks), h


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--base", required=True); ap.add_argument("--marca", default="JOHN HOLDEN"); ap.add_argument("--salida", default=None)
    a = ap.parse_args()
    t0 = time.time(); res = cargar(a.base); print(f"[motor] {os.path.basename(a.base)} en {time.time()-t0:.1f}s")
    try:
        import venta_perdida_semanal as vps
        df_vp = vps.venta_perdida_semana().get("detalle")
    except Exception as e:
        df_vp = None; print(f"[df_vp] no disponible: {e}")
    corte = str(etl.fecha_corte_desde_nombre(os.path.basename(a.base)) or "s-f").replace("/", ".")
    marcas = rm.marcas_reporte(res["cobertura"]) if a.marca.upper() == "ALL" else [a.marca]
    todo_ok = True
    for m in marcas:
        if rp.slice_marca(res["cobertura"], m).empty:
            print(f"\n== {m}: sin filas en la base"); continue
        df, ok, h = checksum_marca(res, m, df_vp, str(corte), a.salida)
        todo_ok &= ok
        print(f"\n== {m} · B1 {h['b1']['n_skus']} / B2a {h['b2a']['n_skus']} / B2b {h['b2b']['n_skus']} / B3 {h['b3']['n_skus']} · {'OK' if ok else 'FALLA'}")
        print(df.to_string(index=False))
    print("\nRESULTADO:", "✅ todos los cuadres OK" if todo_ok else "❌ hay cuadres que fallan")
    sys.exit(0 if todo_ok else 1)


if __name__ == "__main__":
    main()
